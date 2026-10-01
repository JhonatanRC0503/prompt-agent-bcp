"""Extracción del texto de los documentos que adjunta el analista.

Se extrae una vez aquí y se inyecta en los tres especialistas, en vez de subir el mismo archivo a
tres sandboxes de code interpreter distintos. De paso, ese texto alimenta al intake: es de donde
sale el RUC cuando el analista adjunta un reporte sin escribirlo.

Para PDFs se usa Document Intelligence (`prebuilt-layout`, salida Markdown) si está configurado.
Importa porque estos documentos son tablas —estados financieros, líneas de crédito, calificaciones—
y `pypdf` las devuelve como texto plano desordenado, mientras que layout las devuelve como tablas
Markdown. Además hace OCR, así que cubre los PDFs escaneados.
"""

from __future__ import annotations

import io
import logging

from config import AJUSTES
from sesiones import Adjunto

log = logging.getLogger(__name__)

# Por debajo de esto el PDF no tiene capa de texto: es un escaneo.
MIN_CHARS_PDF = 120


async def _texto_document_intelligence(datos: bytes) -> str:
    if not AJUSTES.document_intelligence_endpoint:
        return ""
    try:
        from azure.ai.documentintelligence.aio import DocumentIntelligenceClient
        from azure.ai.documentintelligence.models import DocumentContentFormat
        from azure.identity.aio import DefaultAzureCredential
    except ImportError:
        log.warning("Falta azure-ai-documentintelligence; se usará pypdf.")
        return ""

    async with DefaultAzureCredential() as credencial:
        async with DocumentIntelligenceClient(
            endpoint=AJUSTES.document_intelligence_endpoint, credential=credencial
        ) as cliente:
            operacion = await cliente.begin_analyze_document(
                "prebuilt-layout",
                body=datos,
                output_content_format=DocumentContentFormat.MARKDOWN,
            )
            resultado = await operacion.result()
            return (resultado.content or "").strip()


def _texto_pypdf(datos: bytes) -> str:
    from pypdf import PdfReader

    partes = []
    for pagina in PdfReader(io.BytesIO(datos)).pages:
        try:
            partes.append(pagina.extract_text() or "")
        except Exception:
            log.warning("Página ilegible en el PDF", exc_info=True)
    return "\n".join(partes).strip()


def _texto_xlsx(datos: bytes) -> str:
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
    partes = []
    for hoja in libro.worksheets:
        partes.append(f"### Hoja: {hoja.title}")
        for fila in hoja.iter_rows(values_only=True):
            celdas = [str(c) for c in fila if c is not None]
            if celdas:
                partes.append(" | ".join(celdas))
    libro.close()
    return "\n".join(partes).strip()


def _texto_docx(datos: bytes) -> str:
    import docx

    documento = docx.Document(io.BytesIO(datos))
    partes = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    return "\n".join(partes).strip()


async def _extraer_pdf(adjunto: Adjunto, datos: bytes) -> None:
    try:
        texto = await _texto_document_intelligence(datos)
        if texto:
            adjunto.texto, adjunto.extraido_con = texto, "document-intelligence (prebuilt-layout)"
            return
    except Exception:
        log.warning("Document Intelligence falló; se intenta con pypdf", exc_info=True)

    adjunto.texto = _texto_pypdf(datos)
    adjunto.extraido_con = "pypdf"
    if len(adjunto.texto) < MIN_CHARS_PDF:
        # Escaneado y sin DI disponible. Se declara en vez de pasarlo como documento vacío.
        adjunto.extraido_con = "ninguno (PDF escaneado, sin OCR disponible)"


async def extraer(nombre: str, datos: bytes) -> Adjunto:
    adjunto = Adjunto(nombre=nombre, tipo=nombre.rsplit(".", 1)[-1].lower(), bytes=len(datos))
    minuscula = nombre.lower()

    try:
        if minuscula.endswith(".pdf"):
            await _extraer_pdf(adjunto, datos)
        elif minuscula.endswith((".png", ".jpg", ".jpeg", ".tiff", ".bmp")):
            adjunto.texto = await _texto_document_intelligence(datos)
            adjunto.extraido_con = (
                "document-intelligence (prebuilt-layout)"
                if adjunto.texto
                else "ninguno (imagen sin OCR disponible)"
            )
        elif minuscula.endswith((".xlsx", ".xlsm")):
            adjunto.texto = _texto_xlsx(datos)
            adjunto.extraido_con = "openpyxl"
        elif minuscula.endswith(".docx"):
            adjunto.texto = _texto_docx(datos)
            adjunto.extraido_con = "python-docx"
        else:
            adjunto.texto = datos.decode("utf-8", errors="ignore").strip()
            adjunto.extraido_con = "texto plano"
    except Exception as error:
        log.exception("Fallo extrayendo %s", nombre)
        adjunto.extraido_con = f"error: {error}"

    if len(adjunto.texto) > AJUSTES.max_chars_adjunto:
        adjunto.texto = adjunto.texto[: AJUSTES.max_chars_adjunto] + "\n[...documento truncado...]"
    return adjunto


def como_contexto(adjuntos: list[Adjunto]) -> str:
    """El bloque de texto de los adjuntos que se inyecta en el prompt de los agentes."""
    if not adjuntos:
        return ""
    bloques = []
    for a in adjuntos:
        if a.texto:
            bloques.append(f"--- Documento adjunto: {a.nombre} ---\n{a.texto}")
        else:
            bloques.append(
                f"--- Documento adjunto: {a.nombre} ---\n"
                f"[No se pudo extraer texto ({a.extraido_con}). No asumas su contenido.]"
            )
    return "\n\n".join(bloques)
